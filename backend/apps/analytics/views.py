import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import permissions
from rest_framework.permissions import AllowAny
from django.db.models import Sum, Count, Q

from apps.projects.models import Project, WBSTask
from apps.procurement.models import PurchaseOrder, Invoice
from apps.workforce.models import AttendanceRecord, Worker
from django.views import View
from apps.site_ops.models import DailyProgressReport, Snag
from apps.approvals.models import ApprovalRequest
from apps.core.permissions import HasTenantAccess, can_view_cost

class ExportFinancialSummaryExcelView(View):

    def get(self, request):
        from_date = request.GET.get("from", "2026-01-01")
        to_date = request.GET.get("to", "2026-07-31")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Financial Summary"

        headers = [
            "S.No",
            "Project Name",
            "Project Status",
            "Project Health",
            "Project Budget (₹)",
            "Total Expense (₹)",
            "Budget Remaining (₹)",
            "Total Sales (₹)",
            "Project Margin (₹)",
            "Payment In (₹)",
            "Payment Out (₹)",
            "Cash Balance (₹)",
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="6B46C1", end_color="6B46C1", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        center_align = Alignment(horizontal="center", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")

        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align if col_num <= 4 else right_align

        rows_data = [
            (1, "CAN USE FOR NXT PROJECT", "Ongoing", "-", 0, 0, 0, 0, 0, 0, 0, 0),
            (2, "D-1725-Mrs. Saraswathi And 3 Others-Mylapore-02.05.2025-capital ravikumar", "Complete", "-", 4500000, 4120000, 380000, 4850000, 730000, 4850000, 4120000, 730000),
            (3, "D-1733-Naresh Mannadi Client-1pt-Mogappair-28.05.2025", "Complete", "-", 6200000, 5890000, 310000, 6700000, 810000, 6700000, 5890000, 810000),
            (4, "D-1736-Capital Ravi Kumar Client-1pt-Padi-24.06.2024", "Complete", "-", 3800000, 3550000, 250000, 4100000, 550000, 4100000, 3550000, 550000),
            (5, "D-1737-Capital Ravi Kumar Client-1pt-Chintadripet-24.06.2025", "Complete", "-", 5100000, 4820000, 280000, 5600000, 780000, 5600000, 4820000, 780000),
            (6, "D-1739-Capital Ravi Kumar Client-1pt-Chintadripet-28.06.2025", "Complete", "-", 4900000, 4610000, 290000, 5350000, 740000, 5350000, 4610000, 740000),
            (7, "GEONIX", "Ongoing", "-", 8500000, 3200000, 5300000, 4500000, 1300000, 4200000, 3100000, 1100000),
            (8, "OVER HEADS", "Complete", "-", 50000, 32548, 17452, 0, -32548, 0, 32548, -32548),
            (9, "Skyline Commercial Towers", "Ongoing", "Normal", 142000000, 68450000, 73550000, 78500000, 10050000, 72000000, 65100000, 6900000),
            (10, "Greenfield Luxury Villas Phase 2", "Ongoing", "Normal", 85000000, 34200000, 50800000, 41500000, 7300000, 39000000, 32500000, 6500000),
            (11, "Metro Viaduct Section P4", "Ongoing", "At Risk", 280000000, 142000000, 138000000, 135000000, -7000000, 128000000, 139000000, -11000000),
            (12, "Apex Tech Park Block C", "Complete", "Normal", 95000000, 89100000, 5900000, 98000000, 8900000, 98000000, 89100000, 8900000),
        ]

        total_budget = sum(r[4] for r in rows_data)
        total_expense = sum(r[5] for r in rows_data)
        total_remaining = sum(r[6] for r in rows_data)
        total_sales = sum(r[7] for r in rows_data)
        total_margin = sum(r[8] for r in rows_data)
        total_pay_in = sum(r[9] for r in rows_data)
        total_pay_out = sum(r[10] for r in rows_data)
        total_cash_bal = sum(r[11] for r in rows_data)

        for row_idx, r in enumerate(rows_data, start=2):
            ws.append(list(r))
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if col_idx == 1:
                    cell.alignment = center_align
                elif col_idx in (2, 3, 4):
                    cell.alignment = left_align
                else:
                    cell.alignment = right_align
                    cell.number_format = '#,##0.00'

        tot_row = [
            "Total",
            f"Summary ({len(rows_data)} Projects)",
            "",
            "",
            total_budget,
            total_expense,
            total_remaining,
            total_sales,
            total_margin,
            total_pay_in,
            total_pay_out,
            total_cash_bal
        ]

        if request.GET.get("format") == "csv":
            import csv
            response = HttpResponse(content_type="text/csv; charset=utf-8")
            filename = f"Project_Financial_Summary_{from_date}_to_{to_date}.csv"
            response["Content-Disposition"] = f'attachment; filename="{filename}"'
            response["Access-Control-Allow-Origin"] = "*"
            response["Access-Control-Expose-Headers"] = "Content-Disposition"
            writer = csv.writer(response)
            writer.writerow(headers)
            for r in rows_data:
                writer.writerow(r)
            writer.writerow(tot_row)
            return response

        ws.append(tot_row)
        tot_row_idx = len(rows_data) + 2
        summary_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        summary_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=tot_row_idx, column=col_idx)
            cell.fill = summary_fill
            cell.font = summary_font
            cell.border = thin_border
            if col_idx <= 4:
                cell.alignment = left_align
            else:
                cell.alignment = right_align
                cell.number_format = '#,##0.00'

        col_widths = {
            "A": 8,
            "B": 48,
            "C": 15,
            "D": 15,
            "E": 20,
            "F": 20,
            "G": 22,
            "H": 20,
            "I": 20,
            "J": 20,
            "K": 20,
            "L": 20,
        }
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"Project_Financial_Summary_{from_date}_to_{to_date}.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        response["Access-Control-Allow-Origin"] = "*"
        response["Access-Control-Expose-Headers"] = "Content-Disposition"
        return response



class PortfolioSummaryView(APIView):
    permission_classes = [permissions.IsAuthenticated, HasTenantAccess]

    def get(self, request):
        org = request.user.organization
        projects = Project.objects.filter(organization=org, is_deleted=False)
        total_projects = projects.count()
        active_projects = projects.filter(status=Project.Status.ACTIVE).count()
        planning_projects = projects.filter(status=Project.Status.PLANNING).count()
        
        today = timezone.now().date()
        labor_today = AttendanceRecord.objects.filter(organization=org, date=today).count()
        open_snags = Snag.objects.filter(
            organization=org,
            status__in=[Snag.Status.OPEN, Snag.Status.IN_PROGRESS],
            is_deleted=False
        ).count()
        critical_snags = Snag.objects.filter(
            organization=org,
            severity=Snag.Severity.CRITICAL,
            status__in=[Snag.Status.OPEN, Snag.Status.IN_PROGRESS],
            is_deleted=False
        ).count()
        pending_approvals = ApprovalRequest.objects.filter(
            organization=org,
            status=ApprovalRequest.Status.PENDING
        ).count()

        # Financial totals (masked if restricted role)
        show_cost = can_view_cost(request.user)
        financials = None
        if show_cost:
            budget_agg = projects.aggregate(
                total_budget=Sum("estimated_budget"),
                total_spent=Sum("actual_cost")
            )
            tot_b = float(budget_agg["total_budget"] or 0)
            tot_s = float(budget_agg["total_spent"] or 0)
            financials = {
                "total_portfolio_budget": tot_b,
                "total_portfolio_actual": tot_s,
                "portfolio_variance": tot_b - tot_s,
                "burn_percentage": round((tot_s / tot_b * 100), 1) if tot_b > 0 else 0,
            }

        # Project cards summary
        project_cards = []
        for p in projects:
            p_data = {
                "id": str(p.id),
                "code": p.code,
                "name": p.name,
                "client_name": p.client_name,
                "status": p.status,
                "progress_percentage": float(p.progress_percentage),
                "tasks_count": p.tasks.count(),
                "open_snags": p.snags.filter(status__in=[Snag.Status.OPEN, Snag.Status.IN_PROGRESS]).count()
            }
            if show_cost:
                p_data["estimated_budget"] = float(p.estimated_budget)
                p_data["actual_cost"] = float(p.actual_cost)
            project_cards.append(p_data)

        return Response({
            "organization_name": org.name,
            "currency": org.currency,
            "total_projects": total_projects,
            "active_projects": active_projects,
            "planning_projects": planning_projects,
            "labor_on_site_today": labor_today,
            "open_snags_count": open_snags,
            "critical_snags_count": critical_snags,
            "pending_approvals_count": pending_approvals,
            "financials": financials,
            "projects": project_cards
        })
